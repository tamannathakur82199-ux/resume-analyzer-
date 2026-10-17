"""Builds a resume in the required ATS-friendly format (DOCX + PDF).

A *profile* dict is the single source of truth:
    name, email, phone, location, linkedin, github, summary (str),
    skills {category: [items]}, experience / projects [{heading, date, sub, bullets}],
    education [{heading, date, sub, bullets}], certifications [str], achievements [str]
Anything that could not be found is filled with a clearly marked [placeholder]
(shown in red italics) - the builder never invents facts about the candidate.
"""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from fpdf import FPDF

from resume_parser import BULLET_RE, DATE_RANGE_RE, find_skills
from skills_db import SECTION_TITLES, SKILL_CATEGORY

NAVY = (31, 56, 100)
PH_RED = (192, 57, 43)


def PH(text: str) -> str:
    return f"[{text}]"


def is_ph(text: str) -> bool:
    return isinstance(text, str) and text.startswith("[") and text.endswith("]")


# --------------------------------------------------------------------------- text -> structure
def _split_date(text: str):
    m = DATE_RANGE_RE.search(text)
    if not m:
        return text.strip(), ""
    rest = (text[:m.start()] + " " + text[m.end():]).strip(" |,-–—()\t")
    return re.sub(r"\s{2,}", " ", rest), m.group(0).strip()


def lines_to_entries(lines, flat: bool = False):
    """Turn raw section lines into [{heading, date, sub, bullets}] entries."""
    entries, cur = [], None

    def new(heading="", bullets=None):
        nonlocal cur
        h, d = _split_date(heading) if heading else ("", "")
        cur = {"heading": h, "date": d, "sub": "", "bullets": list(bullets or [])}
        entries.append(cur)

    for raw in lines:
        line = (raw or "").strip()
        text = BULLET_RE.sub("", line).strip()
        if not text:
            continue
        if BULLET_RE.match(line):
            if cur is None:
                new()
            cur["bullets"].append(text)
            continue
        has_date = bool(DATE_RANGE_RE.search(text))
        short = len(text) <= 110
        if flat:
            if cur and re.match(r"(cgpa|gpa|percentage|marks|grade|board|score)\b", text, re.I):
                cur["sub"] = f"{cur['sub']} | {text}" if cur["sub"] else text
            else:
                new(text)
        elif cur is None:
            new(text) if short else new("", [text])
        elif cur["bullets"]:
            if short and (has_date or (len(text.split()) <= 12 and not text.endswith("."))):
                new(text)
            else:
                cur["bullets"].append(text)
        else:
            if short and not cur["sub"]:
                h, d = _split_date(text)
                if d and not cur["date"]:
                    cur["date"] = d
                cur["sub"] = h
            else:
                cur["bullets"].append(text)
    return [e for e in entries if e["heading"] or e["bullets"]]


def lines_to_list(lines):
    return [BULLET_RE.sub("", l).strip() for l in lines if BULLET_RE.sub("", l).strip()]


def skills_from_lines(lines, extra_text: str = ""):
    """Keep explicit 'Category: a, b' groups; auto-group the rest by known category."""
    grouped, loose = {}, []
    for line in lines:
        line = BULLET_RE.sub("", line).strip()
        if not line:
            continue
        if ":" in line and len(line.split(":", 1)[0]) <= 30:
            cat, rest = line.split(":", 1)
            items = [i.strip(" .") for i in re.split(r"[,|;•●▪]+", rest) if i.strip(" .")]
            if items:
                grouped.setdefault(cat.strip().title(), []).extend(items)
            continue
        loose += [i.strip(" .-") for i in re.split(r"[,|;•●▪]+|\s{2,}", line) if 1 < len(i.strip(" .-")) <= 40]
    present = {i.lower() for v in grouped.values() for i in v}
    auto = {}

    def put(item):
        canon = list(find_skills(item, fuzzy=False))
        name = canon[0] if canon else item
        if name.lower() in present or item.lower() in present:
            return
        present.add(name.lower())
        auto.setdefault(SKILL_CATEGORY.get(name, "Other"), []).append(name)

    for it in loose:
        put(it)
    for s in find_skills(extra_text, fuzzy=False):  # skills actually used in projects/experience
        put(s)
    for cat, items in auto.items():
        grouped.setdefault(cat, []).extend(items)
    return grouped


# --------------------------------------------------------------------------- profile builders
DEGREE_RE = re.compile(r"\b(MCA|BCA|B\.?\s?Tech|M\.?\s?Tech|B\.?E\.?|B\.?Sc|M\.?Sc|MBA|BBA|B\.?Com|M\.?Com|Ph\.?D)\b", re.I)


def build_profile(parsed: dict) -> dict:
    sec = parsed["sections"]
    has_skill_section = bool(sec.get("skills"))
    # Skills section present -> trust it. Otherwise (unstructured resume) take every skill found in the text.
    extra = "" if has_skill_section else re.sub(r"\S*(?:github|linkedin)\.com\S*", " ", parsed["raw_text"])
    education = lines_to_entries(sec.get("education", []), flat=True)
    if not education:
        m = DEGREE_RE.search(parsed["raw_text"])
        if m:  # degree is known, the rest is left for the candidate to fill
            education = [{"heading": m.group(1).upper().replace(" ", ""), "date": PH("Year - Year"),
                          "sub": PH("College / University, CGPA"), "bullets": []}]
    return finalize_profile({
        "name": parsed["name"], "email": parsed["email"], "phone": parsed["phone"],
        "linkedin": parsed["linkedin"], "github": parsed["github"], "location": "",
        "summary": " ".join(sec.get("summary", [])).strip(),
        "skills": skills_from_lines(sec.get("skills", []), extra),
        "experience": lines_to_entries(sec.get("experience", [])),
        "projects": lines_to_entries(sec.get("projects", [])),
        "education": education,
        "certifications": lines_to_list(sec.get("certifications", [])),
        "achievements": lines_to_list(sec.get("achievements", [])),
    })


def profile_from_fields(f: dict) -> dict:
    """Profile from plain-text form fields (used by the 'Build Resume' tab)."""
    def lines(key):
        return [l for l in (f.get(key) or "").split("\n") if l.strip()]
    body = "" if (f.get("skills") or "").strip() else \
        (f.get("experience") or "") + "\n" + (f.get("projects") or "")
    return finalize_profile({
        "name": (f.get("name") or "").strip(), "email": (f.get("email") or "").strip(),
        "phone": (f.get("phone") or "").strip(), "location": (f.get("location") or "").strip(),
        "linkedin": (f.get("linkedin") or "").strip(), "github": (f.get("github") or "").strip(),
        "summary": (f.get("summary") or "").strip(),
        "skills": skills_from_lines(lines("skills"), body),
        "experience": lines_to_entries(lines("experience")),
        "projects": lines_to_entries(lines("projects")),
        "education": lines_to_entries(lines("education"), flat=True),
        "certifications": lines_to_list(lines("certifications")),
        "achievements": lines_to_list(lines("achievements")),
    })


def finalize_profile(p: dict) -> dict:
    """Fill every missing mandatory part with a visible placeholder."""
    p = dict(p)
    p["name"] = p.get("name") or PH("Your Full Name")
    p["email"] = p.get("email") or PH("your.email@example.com")
    p["phone"] = p.get("phone") or PH("10-digit mobile number")
    p["summary"] = p.get("summary") or PH(
        "Write 2-3 lines: who you are, your top skills and the role you want")
    if not p.get("skills"):
        p["skills"] = {"Technical Skills": [PH("Add skills grouped as Languages, Frameworks, Databases, Tools")]}
    if not p.get("education"):
        p["education"] = [{"heading": PH("Degree, College / University"), "date": PH("Year - Year"),
                           "sub": PH("CGPA / Percentage"), "bullets": []}]
    if not p.get("experience") and not p.get("projects"):
        p["projects"] = [{"heading": PH("Project name | Tech stack"), "date": "", "sub": "",
                          "bullets": [PH("Built ... using ... that achieved ... (add a number / result)")]}]
    return p


# --------------------------------------------------------------------------- shared layout
def profile_to_blocks(p: dict):
    contact = [p["email"], p["phone"]] + [p.get(k) for k in ("location", "linkedin", "github") if p.get(k)]
    blocks = [("name", p["name"]), ("contact", contact),
              ("section", SECTION_TITLES["summary"]), ("para", p["summary"]),
              ("section", SECTION_TITLES["skills"])]
    for cat, items in p["skills"].items():
        blocks.append(("skill", cat, ", ".join(items)))
    for key in ("experience", "projects", "education"):
        if p.get(key):
            blocks.append(("section", SECTION_TITLES[key]))
            for e in p[key]:
                blocks.append(("entry", e["heading"], e["date"], e["sub"]))
                blocks += [("bullet", b) for b in e["bullets"]]
    for key in ("certifications", "achievements"):
        if p.get(key):
            blocks.append(("section", SECTION_TITLES[key]))
            blocks += [("bullet", t) for t in p[key]]
    return blocks


# --------------------------------------------------------------------------- DOCX
def _docx_run(par, text, bold=False, italic=False, size=None, color=None):
    r = par.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if is_ph(text):
        r.italic, color = True, PH_RED
    if color:
        r.font.color.rgb = RGBColor(*color)
    return r


def _docx_rule(par):
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "1F3864")):
        bottom.set(qn(f"w:{k}"), v)
    pbdr.append(bottom)
    par._p.get_or_add_pPr().append(pbdr)


def build_resume_docx(profile: dict, path: str | None = None) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(17)
    sec.top_margin = sec.bottom_margin = Mm(14)
    usable = Mm(210 - 34)
    for style_name in ("Normal", "List Bullet"):
        st = doc.styles[style_name]
        st.font.name, st.font.size = "Calibri", Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    for block in profile_to_blocks(profile):
        kind = block[0]
        if kind == "name":
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(1)
            _docx_run(par, block[1], bold=True, size=20, color=NAVY)
        elif kind == "contact":
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(4)
            for i, item in enumerate(block[1]):
                if i:
                    _docx_run(par, "  |  ", size=10)
                _docx_run(par, item, size=10)
        elif kind == "section":
            par = doc.add_paragraph()
            par.paragraph_format.space_before, par.paragraph_format.space_after = Pt(8), Pt(3)
            par.paragraph_format.keep_with_next = True
            _docx_run(par, block[1].upper(), bold=True, size=11, color=NAVY)
            _docx_rule(par)
        elif kind == "para":
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(2)
            _docx_run(par, block[1])
        elif kind == "skill":
            par = doc.add_paragraph()
            par.paragraph_format.space_after = Pt(1)
            _docx_run(par, block[1] + ": ", bold=True)
            _docx_run(par, block[2])
        elif kind == "entry":
            par = doc.add_paragraph()
            par.paragraph_format.space_before, par.paragraph_format.space_after = Pt(4), Pt(0)
            par.paragraph_format.keep_with_next = True
            par.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
            _docx_run(par, block[1], bold=True)
            if block[2]:
                _docx_run(par, "\t")
                _docx_run(par, block[2], italic=True, size=10)
            if block[3]:
                sub = doc.add_paragraph()
                sub.paragraph_format.space_after = Pt(1)
                _docx_run(sub, block[3], italic=True, size=10, color=(80, 80, 80))
        elif kind == "bullet":
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(1)
            _docx_run(par, block[1])
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    if path:
        with open(path, "wb") as fh:
            fh.write(data)
    return data


# --------------------------------------------------------------------------- PDF
_LATIN = {"–": "-", "—": "-", "•": "-", "’": "'", "‘": "'", "“": '"', "”": '"',
          "…": "...", " ": " ", "₹": "Rs.", "→": "->", "✓": "-"}


def _latin(s: str) -> str:
    for k, v in _LATIN.items():
        s = s.replace(k, v)
    return s.encode("latin-1", "replace").decode("latin-1")


def build_resume_pdf(profile: dict, path: str | None = None) -> bytes:
    pdf = FPDF(unit="mm", format="A4")
    pdf.set_margins(16, 13, 16)
    pdf.set_auto_page_break(True, 13)
    pdf.add_page()
    W = pdf.w - 32

    def color_for(text, default=(0, 0, 0)):
        pdf.set_text_color(*(PH_RED if is_ph(text) else default))

    def font(style="", size=10, text=""):
        pdf.set_font("Helvetica", ("I" if is_ph(text) and "I" not in style else "") + style, size)

    for block in profile_to_blocks(profile):
        kind = block[0]
        if kind == "name":
            font("B", 20, block[1]); color_for(block[1], NAVY)
            pdf.cell(0, 9, _latin(block[1]), align="C", new_x="LMARGIN", new_y="NEXT")
        elif kind == "contact":
            items, sep = [_latin(i) for i in block[1]], "   |   "
            font("", 9.5)
            total = sum(pdf.get_string_width(i) for i in items) + pdf.get_string_width(sep) * (len(items) - 1)
            pdf.set_x(max(pdf.l_margin, (pdf.w - total) / 2))
            for n, item in enumerate(items):
                if n:
                    pdf.set_text_color(90, 90, 90); font("", 9.5); pdf.write(5, sep)
                font("", 9.5, item); color_for(item, (40, 40, 40)); pdf.write(5, item)
            pdf.ln(7)
        elif kind == "section":
            pdf.ln(2.5)
            font("B", 11); pdf.set_text_color(*NAVY)
            pdf.cell(0, 6, _latin(block[1].upper()), new_x="LMARGIN", new_y="NEXT")
            y = pdf.get_y()
            pdf.set_draw_color(*NAVY); pdf.set_line_width(0.3)
            pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
            pdf.ln(1.5)
        elif kind == "para":
            font("", 10, block[1]); color_for(block[1])
            pdf.multi_cell(0, 5, _latin(block[1]), new_x="LMARGIN", new_y="NEXT")
        elif kind == "skill":
            font("B", 10); pdf.set_text_color(0, 0, 0)
            label = _latin(block[1]) + ": "
            lw = pdf.get_string_width(label) + 0.5
            y0 = pdf.get_y()
            pdf.cell(lw, 5, label)
            font("", 10, block[2]); color_for(block[2])
            pdf.multi_cell(W - lw, 5, _latin(block[2]), new_x="LMARGIN", new_y="NEXT")
            if pdf.get_y() == y0:  # safety
                pdf.ln(5)
        elif kind == "entry":
            _, heading, date, sub = block
            pdf.ln(1.2)
            font("B", 10.5, heading); color_for(heading)
            date_w = 0
            if date:
                font("I", 9.5, date); date_w = pdf.get_string_width(_latin(date)) + 2
            font("B", 10.5, heading)
            if pdf.get_string_width(_latin(heading)) <= W - date_w - 2:
                pdf.cell(W - date_w, 5.5, _latin(heading), new_x="RIGHT", new_y="TOP")
                if date:
                    font("I", 9.5, date); color_for(date, (60, 60, 60))
                    pdf.cell(date_w, 5.5, _latin(date), align="R", new_x="LMARGIN", new_y="NEXT")
                else:
                    pdf.ln(5.5)
            else:
                pdf.multi_cell(0, 5.5, _latin(heading), new_x="LMARGIN", new_y="NEXT")
                if date:
                    font("I", 9.5, date); color_for(date, (60, 60, 60))
                    pdf.cell(0, 5, _latin(date), new_x="LMARGIN", new_y="NEXT")
            if sub:
                font("I", 9.5, sub); color_for(sub, (80, 80, 80))
                pdf.multi_cell(0, 4.8, _latin(sub), new_x="LMARGIN", new_y="NEXT")
        elif kind == "bullet":
            font("", 10, block[1]); color_for(block[1])
            if pdf.get_y() > pdf.h - 22:
                pdf.add_page()
            x0, y0 = pdf.l_margin, pdf.get_y()
            pdf.set_fill_color(*(PH_RED if is_ph(block[1]) else (40, 40, 40)))
            pdf.ellipse(x0 + 2.2, y0 + 1.9, 1.2, 1.2, style="F")
            pdf.set_x(x0 + 6)
            pdf.multi_cell(W - 6, 4.8, _latin(block[1]), new_x="LMARGIN", new_y="NEXT")
    data = bytes(pdf.output())
    if path:
        with open(path, "wb") as fh:
            fh.write(data)
    return data
