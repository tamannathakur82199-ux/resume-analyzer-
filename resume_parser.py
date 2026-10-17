"""Resume parsing: text extraction (PDF / DOCX / TXT) and structured field detection."""
from __future__ import annotations

import datetime as _dt
import io
import os
import re

from skills_db import (
    ALL_SKILLS, SECTION_ALIASES, SKILL_ALIASES, is_case_sensitive,
)

try:  # fast typo matching; optional - falls back to the standard library if unavailable/blocked
    from rapidfuzz import fuzz, process

    def _fuzzy_hit(key, tokens):
        return process.extractOne(key, tokens, scorer=fuzz.ratio, score_cutoff=88)
except Exception:  # ImportError, or a DLL blocked by Windows policy
    import difflib

    def _fuzzy_hit(key, tokens):
        for tok in tokens:
            if (abs(len(tok) - len(key)) <= 2 and tok[:1] == key[:1]
                    and difflib.SequenceMatcher(None, key, tok).ratio() >= 0.88):
                return tok
        return None

# \uf0b7 / \uf0a7 / \uf0d8 / \uf0fc / \uf076 = bullet glyphs from Symbol / Wingdings fonts (Word and LibreOffice PDFs)
BULLET_CHARS = "•●▪■◦◆►➢➤✓✔–—-*·○\uf0b7\uf0a7\uf0d8\uf0fc\uf076\uf0a8\uf06e"
BULLET_RE = re.compile(r"^\s*[" + re.escape(BULLET_CHARS) + r"]+\s*")

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}

_DATE = (r"(?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?,?\s*\d{4}"
         r"|\d{1,2}\s*[/\-]\s*\d{4}|\d{4})")
_END = r"(?:present|current|now|till\s+date|ongoing|" + _DATE + r")"
DATE_RANGE_RE = re.compile(r"(" + _DATE + r")\s*(?:-|–|—|to)\s*(" + _END + r")", re.I)


# --------------------------------------------------------------------------- text extraction
def extract_text(source, filename: str | None = None):
    """Return (text, meta). `source` may be a path, bytes or a file-like object."""
    if isinstance(source, (str, os.PathLike)):
        filename = filename or os.fspath(source)
        with open(source, "rb") as fh:
            data = fh.read()
    elif hasattr(source, "read"):
        data = source.read()
        filename = filename or getattr(source, "name", "resume")
    else:
        data = bytes(source)
    ext = os.path.splitext(filename or "")[1].lower()
    meta = {"format": ext.lstrip(".") or "txt", "pages": 1}

    if ext == ".pdf":
        import pdfplumber
        parts = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            meta["pages"] = len(pdf.pages)
            for page in pdf.pages:
                parts.append(page.extract_text() or "")
        text = "\n".join(parts)
    elif ext in (".docx", ".dotx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        lines = []
        for p in doc.paragraphs:
            pPr = p._p.pPr
            is_list = (pPr is not None and pPr.numPr is not None) or \
                (p.style is not None and p.style.name.lower().startswith("list"))
            lines.append(("• " + p.text) if is_list and p.text.strip() else p.text)
        meta["has_tables"] = bool(doc.tables)
        for table in doc.tables:
            for row in table.rows:
                seen = []
                for cell in row.cells:
                    if cell.text not in seen:
                        seen.append(cell.text)
                lines.append("  ".join(seen))
        text = "\n".join(lines)
        meta["pages"] = max(1, round(len(text.split()) / 520))
    else:
        text = data.decode("utf-8", errors="replace")
        meta["pages"] = max(1, round(len(text.split()) / 520))
    text = text.replace("\r", "\n").replace(" ", " ").replace("\t", "  ")
    return text, meta


# --------------------------------------------------------------------------- skills
_SKILL_PATTERNS = None


def _compile_patterns():
    pairs = {}
    for s in ALL_SKILLS:
        pairs[s] = s
    for alias, canon in SKILL_ALIASES.items():
        pairs[alias] = canon
    compiled = []
    for term, canon in pairs.items():
        flags = 0 if is_case_sensitive(term) else re.I
        tail = r"(?![A-Za-z0-9+#&])" if len(term) == 1 else r"(?![A-Za-z0-9+#])"
        pat = re.compile(r"(?<![A-Za-z0-9+#])" + re.escape(term).replace(r"\ ", r"[\s\-]?") + tail, flags)
        compiled.append((term, canon, pat))
    return compiled


def find_skills(text: str, fuzzy: bool = True):
    """Return {canonical_skill: 'exact' | 'fuzzy'} found in `text`."""
    global _SKILL_PATTERNS
    if _SKILL_PATTERNS is None:
        _SKILL_PATTERNS = _compile_patterns()
    found = {}
    for _term, canon, pat in _SKILL_PATTERNS:
        if canon not in found and pat.search(text):
            found[canon] = "exact"
    if fuzzy:
        low = text.lower()
        words = re.findall(r"[a-z0-9+#.]+", low)
        tokens = {w.strip(".") for w in words if len(w) >= 5}
        tokens |= {a + b for a, b in zip(words, words[1:]) if len(a + b) >= 6}
        for skill in ALL_SKILLS:
            if skill in found:
                continue
            key = re.sub(r"[\s\-.]", "", skill.lower())
            if len(key) < 6 or " " in skill and len(key) < 8:
                continue
            if _fuzzy_hit(key, tokens):
                found[skill] = "fuzzy"
    return found


# --------------------------------------------------------------------------- sections
def _norm_heading(line: str) -> str:
    low = re.sub(r"[^a-zA-Z& ]", " ", line).lower().replace("&", " and ")
    return re.sub(r"\s+", " ", low).strip()


_HEADING_LOOKUP = {}
for _key, _names in SECTION_ALIASES.items():
    for _n in _names:
        _HEADING_LOOKUP[_n] = _key


def heading_key(line: str):
    stripped = line.strip()
    if not stripped or len(stripped) > 45 or "@" in stripped:
        return None
    return _HEADING_LOOKUP.get(_norm_heading(BULLET_RE.sub("", stripped)))


def split_sections(text: str):
    sections = {"header": []}
    current = "header"
    for raw in text.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        key = heading_key(line)
        if key:
            current = key
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line.strip())
    return sections


# --------------------------------------------------------------------------- fields
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/[^\s,;|)]+", re.I)
GITHUB_RE = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[^\s,;|)]+", re.I)
PHONE_RE = re.compile(r"(?<![\d/])\+?\d[\d\s\-().]{8,17}\d(?![\d/])")


def _find_phone(text: str):
    for m in PHONE_RE.finditer(text):
        cand = m.group().strip()
        digits = re.sub(r"\D", "", cand)
        if not 10 <= len(digits) <= 13:
            continue
        if re.fullmatch(r"\d{4}\s*[-–]\s*\d{4}", cand):
            continue
        return cand
    return ""


def _find_name(header_lines, email: str):
    for line in header_lines[:8]:
        clean = line.strip()
        low = clean.lower()
        if any(t in low for t in ("@", "http", "linkedin", "github", "resume", "curriculum", "vitae")):
            continue
        if re.search(r"\d{5,}", clean) or heading_key(clean):
            continue
        words = clean.replace("|", " ").split()
        if 1 <= len(words) <= 4 and all(re.fullmatch(r"[A-Za-z.'\-]+", w) for w in words):
            return clean.title() if clean.isupper() else clean
    if email:
        local = re.split(r"[._\d]+", email.split("@")[0])
        local = [p for p in local if len(p) > 1]
        if local:
            return " ".join(p.title() for p in local[:2])
    return ""


def _to_date(token: str, is_end: bool):
    token = token.strip().lower()
    today = _dt.date.today()
    if re.match(r"(present|current|now|till|ongoing)", token):
        return today.year * 12 + today.month
    m = re.match(r"([a-z]{3})[a-z]*\.?,?\s*(\d{4})", token)
    if m:
        return int(m.group(2)) * 12 + MONTHS[m.group(1)]
    m = re.match(r"(\d{1,2})\s*[/\-]\s*(\d{4})", token)
    if m:
        return int(m.group(2)) * 12 + min(12, max(1, int(m.group(1))))
    m = re.match(r"(\d{4})", token)
    if m:
        return int(m.group(1)) * 12 + 1
    return None


def estimate_experience_years(exp_lines, text: str) -> float:
    spans = []
    for m in DATE_RANGE_RE.finditer("\n".join(exp_lines)):
        a, b = _to_date(m.group(1), False), _to_date(m.group(2), True)
        if a is not None and b is not None and b >= a:
            spans.append((a, b))
    spans.sort()
    merged = []
    for a, b in spans:
        if merged and a <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], b))
        else:
            merged.append((a, b))
    from_dates = sum(b - a for a, b in merged) / 12.0
    stated = 0.0
    for m in re.finditer(r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)(?:\s+of)?\s+(?:[a-z]+\s+){0,3}experience", text, re.I):
        v = float(m.group(1))
        if v <= 40:
            stated = max(stated, v)
    return round(max(from_dates, stated), 1)


def split_skill_items(lines):
    items = []
    for line in lines:
        line = BULLET_RE.sub("", line)
        if ":" in line:
            line = line.split(":", 1)[1]
        for part in re.split(r"[,|;/•●▪]+|\s{2,}", line):
            part = part.strip(" .-")
            if 1 < len(part) <= 40:
                items.append(part)
    seen, out = set(), []
    for it in items:
        if it.lower() not in seen:
            seen.add(it.lower())
            out.append(it)
    return out


def parse_resume(text: str, meta: dict | None = None) -> dict:
    sections = split_sections(text)
    header = sections.get("header", [])
    head_text = "\n".join(header[:10])
    email = (EMAIL_RE.search(head_text) or EMAIL_RE.search(text) or [None])
    email = email.group() if hasattr(email, "group") else ""
    phone = _find_phone(head_text) or _find_phone(text)
    linkedin = (LINKEDIN_RE.search(text) or [None])
    linkedin = linkedin.group() if hasattr(linkedin, "group") else ""
    github = (GITHUB_RE.search(text) or [None])
    github = github.group() if hasattr(github, "group") else ""
    exp_lines = sections.get("experience", [])
    return {
        "name": _find_name(header, email),
        "email": email,
        "phone": phone,
        "linkedin": linkedin,
        "github": github,
        "sections": sections,
        "skills": list(find_skills(re.sub(r"\S*(?:github|linkedin)\.com\S*|https?://\S+", " ", text)).keys()),
        "skills_listed": split_skill_items(sections.get("skills", [])),
        "experience_years": estimate_experience_years(exp_lines, text),
        "word_count": len(text.split()),
        "pages": (meta or {}).get("pages", 1),
        "raw_text": text,
    }
